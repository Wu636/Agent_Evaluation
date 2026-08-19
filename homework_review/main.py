"""
作业批阅后端API服务
提供生成学生答案和批阅评测功能
"""

import asyncio
import json
import os
import subprocess
import sys
import tempfile
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import Depends, FastAPI, File, Form, Header, UploadFile, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from dotenv import load_dotenv

try:
    from .review_job_control import (
        FairUserConcurrencyLimiter,
        ReviewAuthConfigurationError,
        ReviewAuthenticationError,
        SupabaseTokenVerifier,
        extract_bearer_token,
    )
    from .skill_generation_service import (
        encode_file_base64,
        generate_grading_skill_zip,
        generate_student_sample_docx_files,
    )
    from .skill_review_service import (
        build_submission_requirement_from_overview,
        build_skill_score_table,
        compact_skill_report,
        execute_correction_skill,
        get_correction_skill_overview,
        get_correction_skill_report,
        list_correction_skill_models,
        platform_report_url,
        upload_and_prepare_grading_skill,
        upload_student_attachment,
    )
except ImportError:
    from review_job_control import (
        FairUserConcurrencyLimiter,
        ReviewAuthConfigurationError,
        ReviewAuthenticationError,
        SupabaseTokenVerifier,
        extract_bearer_token,
    )
    from skill_generation_service import (
        encode_file_base64,
        generate_grading_skill_zip,
        generate_student_sample_docx_files,
    )
    from skill_review_service import (
        build_submission_requirement_from_overview,
        build_skill_score_table,
        compact_skill_report,
        execute_correction_skill,
        get_correction_skill_overview,
        get_correction_skill_report,
        list_correction_skill_models,
        platform_report_url,
        upload_and_prepare_grading_skill,
        upload_student_attachment,
    )

app = FastAPI(title="作业批阅API", version="1.0.0")

# CORS配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 获取当前脚本目录
SCRIPT_DIR = Path(__file__).parent
GENERATE_SCRIPT = SCRIPT_DIR / "generate_and_review_service.py"
REVIEW_SCRIPT = SCRIPT_DIR / "review_service.py"
env_file = SCRIPT_DIR / ".env"
if not env_file.exists():
    env_file.touch()
load_dotenv(env_file)

MAX_REVIEW_FILES = 150
DEFAULT_REVIEW_CONCURRENCY = 5
MAX_REVIEW_CONCURRENCY = 10
REVIEW_UPLOAD_CHUNK_SIZE = 1024 * 1024
MAX_SKILL_PACKAGE_BYTES = 50 * 1024 * 1024
MAX_SKILL_MATERIAL_FILES = 20
MAX_SKILL_MATERIAL_BYTES = 80 * 1024 * 1024

def env_int(name: str, default: int, *, minimum: int = 1, maximum: int = 100) -> int:
    try:
        value = int(os.getenv(name, str(default)))
    except (TypeError, ValueError):
        value = default
    return min(maximum, max(minimum, value))


MAX_ACTIVE_REVIEW_JOBS_PER_USER = env_int("MAX_ACTIVE_REVIEW_JOBS_PER_USER", 1, maximum=5)
MAX_ACTIVE_TRADITIONAL_REVIEW_JOBS = env_int("MAX_ACTIVE_TRADITIONAL_REVIEW_JOBS", 1, maximum=5)
MAX_GLOBAL_SKILL_ATTEMPTS = env_int("MAX_GLOBAL_SKILL_ATTEMPTS", 10, maximum=50)
MAX_SKILL_ATTEMPTS_PER_USER = env_int("MAX_SKILL_ATTEMPTS_PER_USER", 3, maximum=10)
MAX_GLOBAL_SKILL_UPLOADS = env_int("MAX_GLOBAL_SKILL_UPLOADS", 4, maximum=20)
MAX_SKILL_UPLOADS_PER_USER = env_int("MAX_SKILL_UPLOADS_PER_USER", 2, maximum=10)

# A single Railway process owns the transient task state. Every job is bound to
# a verified Supabase user id, while request-level limiters rotate fairly across
# users instead of locking the entire Skills batch behind one global semaphore.
REVIEW_JOBS: Dict[str, Dict[str, Any]] = {}
REVIEW_JOB_TASKS: set[asyncio.Task] = set()
REVIEW_OWNER_SEMAPHORES: Dict[str, asyncio.Semaphore] = {}
TRADITIONAL_REVIEW_JOB_SEMAPHORE = asyncio.Semaphore(MAX_ACTIVE_TRADITIONAL_REVIEW_JOBS)
SKILL_ATTEMPT_LIMITER = FairUserConcurrencyLimiter(
    MAX_GLOBAL_SKILL_ATTEMPTS,
    MAX_SKILL_ATTEMPTS_PER_USER,
)
SKILL_UPLOAD_LIMITER = FairUserConcurrencyLimiter(
    MAX_GLOBAL_SKILL_UPLOADS,
    MAX_SKILL_UPLOADS_PER_USER,
)
SUPABASE_TOKEN_VERIFIER = SupabaseTokenVerifier.from_env()
SYSTEM_TEMP_ROOT = Path(tempfile.gettempdir()).resolve()
REVIEW_JOBS_ROOT = SYSTEM_TEMP_ROOT / "homework_review_jobs"
REVIEW_JOBS_ROOT.mkdir(parents=True, exist_ok=True)

def clamp_review_concurrency(value: int) -> int:
    return min(MAX_REVIEW_CONCURRENCY, max(1, int(value)))


async def require_review_user(
    authorization_header: Optional[str] = Header(None, alias="Authorization"),
) -> str:
    try:
        token = extract_bearer_token(authorization_header)
        return await asyncio.to_thread(SUPABASE_TOKEN_VERIFIER.verify, token)
    except ReviewAuthenticationError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc
    except ReviewAuthConfigurationError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc


def get_review_job(job_id: str, owner_id: Optional[str] = None) -> Dict[str, Any]:
    job = REVIEW_JOBS.get(job_id)
    if not job or (owner_id is not None and job.get("ownerId") != owner_id):
        raise HTTPException(status_code=404, detail="批阅任务不存在或服务已重启")
    return job


def get_review_owner_semaphore(owner_id: str) -> asyncio.Semaphore:
    semaphore = REVIEW_OWNER_SEMAPHORES.get(owner_id)
    if semaphore is None:
        semaphore = asyncio.Semaphore(MAX_ACTIVE_REVIEW_JOBS_PER_USER)
        REVIEW_OWNER_SEMAPHORES[owner_id] = semaphore
    return semaphore


def resolve_temp_file(path: str) -> Path:
    """Resolve a file while keeping access inside this system's temp directory."""
    resolved = Path(path).resolve()
    try:
        resolved.relative_to(SYSTEM_TEMP_ROOT)
    except ValueError as exc:
        raise HTTPException(status_code=403, detail="只允许访问临时文件") from exc
    return resolved


def resolve_review_job_artifact(job: Dict[str, Any], file_name: str) -> Path:
    """Resolve a result artifact strictly inside this job's output directory."""
    output_root = Path(job["outputRoot"]).resolve()
    requested = Path(file_name)
    candidate = requested.resolve() if requested.is_absolute() else (output_root / requested).resolve()
    try:
        candidate.relative_to(output_root)
    except ValueError as exc:
        raise HTTPException(status_code=403, detail="结果文件路径超出当前批阅任务") from exc
    if not candidate.is_file():
        raise HTTPException(status_code=404, detail="结果文件不存在")
    return candidate


def append_review_job_log(job: Dict[str, Any], message: str, level: str = "info") -> None:
    job["logs"].append({
        "index": len(job["logs"]),
        "message": message,
        "level": level,
    })
    job["updatedAt"] = datetime.now(timezone.utc).isoformat()


def safe_upload_name(filename: str) -> str:
    basename = Path(filename or "file").name
    safe = "".join(
        ch if ch.isalnum() or ch in "-_.()[] " or "\u4e00" <= ch <= "\u9fff" else "_"
        for ch in basename
    ).strip(" .")
    return safe[:180] or "file"


def unique_upload_path(upload_dir: Path, filename: str) -> Path:
    candidate = upload_dir / safe_upload_name(filename)
    if not candidate.exists():
        return candidate

    stem, suffix = candidate.stem, candidate.suffix
    index = 2
    while True:
        candidate = upload_dir / f"{stem}({index}){suffix}"
        if not candidate.exists():
            return candidate
        index += 1


async def terminate_review_job_process(job: Dict[str, Any]) -> None:
    process = job.get("_process")
    if process is None or process.returncode is not None:
        return
    process.terminate()
    try:
        await asyncio.wait_for(process.wait(), timeout=5)
    except asyncio.TimeoutError:
        process.kill()
        await process.wait()


def review_job_summary(job: Dict[str, Any]) -> Dict[str, Any]:
    """Return only the fields needed by the active-job recovery UI."""
    return {
        "jobId": job["jobId"],
        "status": job["status"],
        "engine": job.get("engine") or "traditional",
        "uploadedCount": len(job.get("files") or []),
        "createdAt": job["createdAt"],
        "updatedAt": job["updatedAt"],
    }


async def cancel_review_job_state(job: Dict[str, Any], message: str) -> None:
    """Cancel a job and wait briefly for its per-user execution slot to release."""
    job["cancelRequested"] = True
    await terminate_review_job_process(job)
    job["status"] = "cancelled"
    job["error"] = None
    append_review_job_log(job, message, "warn")

    task = job.get("_task")
    if task is None or task.done() or task is asyncio.current_task():
        return
    task.cancel()
    try:
        await asyncio.wait_for(asyncio.shield(task), timeout=3)
    except (asyncio.CancelledError, asyncio.TimeoutError):
        pass


async def execute_async_review_job(
    job_id: str,
    *,
    authorization: str,
    cookie: str,
    instance_nid: str,
    attempts: int,
    output_format: str,
    max_concurrency: int,
    local_parse: bool,
    llm_api_key: str,
    llm_api_url: str,
    llm_model: str,
    skip_llm_files: Optional[str],
    file_groups: Optional[str],
) -> None:
    job = get_review_job(job_id)
    job["status"] = "running"
    append_review_job_log(
        job,
        f"🚀 后台批阅已启动：{len(job['files'])} 份作业，每份 {attempts} 次，并发 {max_concurrency}",
    )

    env = os.environ.copy()
    env["PYTHONUNBUFFERED"] = "1"
    env["AUTHORIZATION"] = authorization
    env["COOKIE"] = cookie
    env["INSTANCE_NID"] = instance_nid
    env["LLM_API_KEY"] = llm_api_key or os.getenv("LLM_API_KEY", "")
    env["LLM_API_URL"] = llm_api_url or os.getenv("LLM_API_URL", "")
    env["LLM_MODEL"] = llm_model or os.getenv("LLM_MODEL", "")

    cmd = [
        sys.executable, "-u", str(REVIEW_SCRIPT),
        "--inputs", json.dumps(job["files"]),
        "--attempts", str(max(1, attempts)),
        "--output-format", output_format,
        "--output-root", job["outputRoot"],
        "--max-concurrency", str(clamp_review_concurrency(max_concurrency)),
        "--compact-result",
    ]
    if local_parse:
        cmd.append("--local-parse")
    if skip_llm_files:
        cmd.extend(["--skip-llm-files", skip_llm_files])
    if file_groups:
        cmd.extend(["--file-groups", file_groups])

    result_payload: Optional[Dict[str, Any]] = None
    process: Optional[asyncio.subprocess.Process] = None
    try:
        process = await asyncio.create_subprocess_exec(
            *cmd,
            env=env,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
            cwd=str(SCRIPT_DIR),
            limit=32 * 1024 * 1024,
        )
        job["pid"] = process.pid
        job["_process"] = process

        async def read_stdout() -> None:
            nonlocal result_payload
            assert process.stdout is not None
            while True:
                line = await process.stdout.readline()
                if not line:
                    break
                message = line.decode(errors="replace").strip()
                if not message:
                    continue
                if message.startswith("__RESULT__"):
                    result_payload = json.loads(message[len("__RESULT__"):])
                else:
                    append_review_job_log(job, message)

        async def read_stderr() -> None:
            assert process.stderr is not None
            while True:
                line = await process.stderr.readline()
                if not line:
                    break
                message = line.decode(errors="replace").strip()
                if message:
                    append_review_job_log(job, f"⚠️ {message}", "warn")

        await asyncio.gather(read_stdout(), read_stderr())
        return_code = await process.wait()
        if return_code != 0:
            raise RuntimeError(f"批阅进程退出码 {return_code}")
        if not result_payload:
            raise RuntimeError("批阅结果缺少完成标记")

        output_root = Path(job["outputRoot"]).resolve()
        relative_files: List[str] = []
        for item in result_payload.get("output_files", []):
            requested = Path(str(item))
            candidate = requested.resolve() if requested.is_absolute() else (output_root / requested).resolve()
            try:
                relative_name = candidate.relative_to(output_root)
            except ValueError:
                append_review_job_log(job, f"⚠️ 已忽略输出目录之外的结果文件：{item}", "warn")
                continue
            if candidate.is_file():
                relative_files.append(str(relative_name))
        job["result"] = {
            "jobId": job_id,
            "outputFiles": relative_files,
            "summary": result_payload.get("result", {}),
            "scoreTable": result_payload.get("score_table"),
            "downloadBaseUrl": f"/api/homework-review/jobs/{job_id}/artifacts",
        }
        job["status"] = "completed"
        append_review_job_log(job, "🎉 全部批阅完成")
    except asyncio.CancelledError:
        await terminate_review_job_process(job)
        if job.get("status") != "cancelled":
            job["status"] = "cancelled"
            append_review_job_log(job, "⏹️ 批阅任务已取消", "warn")
        raise
    except Exception as exc:
        if process is not None and process.returncode is None:
            await terminate_review_job_process(job)
        job["status"] = "failed"
        job["error"] = str(exc)
        append_review_job_log(job, f"❌ 批阅任务失败：{exc}", "error")
    finally:
        job.pop("pid", None)
        job.pop("_process", None)
        job["updatedAt"] = datetime.now(timezone.utc).isoformat()


async def run_async_review_job(job_id: str, **settings: Any) -> None:
    """Serialize jobs per user and cap only the resource-heavy legacy worker."""
    job = get_review_job(job_id)
    owner_id = job["ownerId"]
    owner_semaphore = get_review_owner_semaphore(owner_id)
    if owner_semaphore.locked():
        append_review_job_log(job, "⏳ 你已有一个批阅任务在执行，本任务在个人队列中等待")
    async with owner_semaphore:
        if TRADITIONAL_REVIEW_JOB_SEMAPHORE.locked():
            append_review_job_log(job, "⏳ 传统批阅工作进程繁忙，本任务等待可用资源")
        async with TRADITIONAL_REVIEW_JOB_SEMAPHORE:
            if job.get("cancelRequested"):
                return
            await execute_async_review_job(job_id, **settings)


SKILL_SUCCESS_STATES = {"SUCCESS"}
SKILL_FAILURE_STATES = {"FAILED", "FAILURE", "ERROR", "CANCELLED", "CANCELED"}


def make_skill_task_id() -> str:
    return f"test-{int(time.time() * 1000)}-{uuid.uuid4().hex[:8]}"


async def execute_skill_attempt(
    *,
    job: Dict[str, Any],
    file_name: str,
    file_index: int,
    attempt_index: int,
    attachment: Dict[str, str],
    authorization: str,
    cookie: str,
    skill_version_id: str,
    skill_nid: str,
    model_name: str,
    submission_requirement: str,
    student_submission: str,
    poll_interval_seconds: int,
) -> Dict[str, Any]:
    task_id = make_skill_task_id()
    report_url = platform_report_url(
        skill_version_id=skill_version_id,
        skill_nid=skill_nid,
        task_id=task_id,
    )
    try:
        await asyncio.to_thread(
            execute_correction_skill,
            skill_version_id=skill_version_id,
            task_id=task_id,
            model_name=model_name,
            submission_requirement=submission_requirement,
            student_submission=student_submission,
            student_attachments=[attachment],
            requirement_attachments=[],
            authorization=authorization,
            cookie=cookie,
        )
        append_review_job_log(
            job,
            f"🧪 「{file_name}」第 {attempt_index} 次 Skills 批阅已启动（{task_id}）",
        )

        started_at = time.monotonic()
        next_wait_log = 60
        while True:
            await asyncio.sleep(poll_interval_seconds)
            response = await asyncio.to_thread(
                get_correction_skill_report,
                task_id,
                authorization,
                cookie,
            )
            skill = ((response.get("data") or {}).get("skill") or {})
            report_status = str(skill.get("reportStatus") or "").upper()
            if report_status in SKILL_SUCCESS_STATES:
                return compact_skill_report(
                    response,
                    file_name=file_name,
                    file_index=file_index,
                    attempt_index=attempt_index,
                    task_id=task_id,
                    report_url=report_url,
                )
            if report_status in SKILL_FAILURE_STATES:
                return {
                    "success": False,
                    "fileName": file_name,
                    "fileIndex": file_index,
                    "attemptIndex": attempt_index,
                    "taskId": task_id,
                    "reportUrl": report_url,
                    "reportStatus": report_status,
                    "error": skill.get("message") or skill.get("error") or f"批阅终态：{report_status}",
                    "items": [],
                    "sections": [],
                }

            elapsed = int(time.monotonic() - started_at)
            if elapsed >= next_wait_log:
                append_review_job_log(
                    job,
                    f"⏳ 「{file_name}」第 {attempt_index} 次仍在批阅，已等待 {elapsed} 秒",
                )
                next_wait_log += 60
    except asyncio.CancelledError:
        raise
    except Exception as exc:
        return {
            "success": False,
            "fileName": file_name,
            "fileIndex": file_index,
            "attemptIndex": attempt_index,
            "taskId": task_id,
            "reportUrl": report_url,
            "reportStatus": "ERROR",
            "error": str(exc),
            "items": [],
            "sections": [],
        }


async def execute_async_skill_review_job(
    job_id: str,
    *,
    authorization: str,
    cookie: str,
    skill_version_id: str,
    skill_nid: str,
    submission_requirement: str,
    student_submission: str,
    model_name: str,
    attempts: int,
    max_concurrency: int,
    poll_interval_seconds: int,
) -> None:
    job = get_review_job(job_id)
    owner_id = job["ownerId"]
    job["status"] = "running"
    job["engine"] = "skill"
    total_runs = len(job["files"]) * attempts
    append_review_job_log(
        job,
        f"🚀 Skills 批量测试已启动：{len(job['files'])} 份作业，每份 {attempts} 次，共 {total_runs} 次",
    )

    concurrency = clamp_review_concurrency(max_concurrency)
    semaphore = asyncio.Semaphore(concurrency)
    uploaded: Dict[int, Dict[str, str]] = {}
    upload_errors: Dict[int, str] = {}

    async def upload_one(file_index: int, file_path: str) -> None:
        file_name = Path(file_path).name
        async with semaphore:
            async with SKILL_UPLOAD_LIMITER.slot(owner_id):
                try:
                    uploaded[file_index] = await asyncio.to_thread(
                        upload_student_attachment,
                        file_path,
                        authorization,
                        cookie,
                    )
                    append_review_job_log(job, f"☁️ 「{file_name}」已上传到智慧树资源服务")
                except Exception as exc:
                    upload_errors[file_index] = str(exc)
                    append_review_job_log(job, f"❌ 「{file_name}」上传失败：{exc}", "error")

    try:
        await asyncio.gather(*[
            upload_one(file_index, file_path)
            for file_index, file_path in enumerate(job["files"])
        ])

        completed_runs = 0
        results: List[Dict[str, Any]] = []

        async def run_one(file_index: int, file_path: str, attempt_index: int) -> None:
            nonlocal completed_runs
            file_name = Path(file_path).name
            if file_index in upload_errors:
                result = {
                    "success": False,
                    "fileName": file_name,
                    "fileIndex": file_index,
                    "attemptIndex": attempt_index,
                    "taskId": "",
                    "reportUrl": "",
                    "reportStatus": "UPLOAD_FAILED",
                    "error": upload_errors[file_index],
                    "items": [],
                    "sections": [],
                }
            else:
                async with semaphore:
                    async with SKILL_ATTEMPT_LIMITER.slot(owner_id):
                        result = await execute_skill_attempt(
                            job=job,
                            file_name=file_name,
                            file_index=file_index,
                            attempt_index=attempt_index,
                            attachment=uploaded[file_index],
                            authorization=authorization,
                            cookie=cookie,
                            skill_version_id=skill_version_id,
                            skill_nid=skill_nid,
                            model_name=model_name,
                            submission_requirement=submission_requirement,
                            student_submission=student_submission,
                            poll_interval_seconds=poll_interval_seconds,
                        )
            results.append(result)
            completed_runs += 1
            if result.get("success"):
                append_review_job_log(
                    job,
                    f"✅ 「{file_name}」第 {attempt_index} 次完成：{result.get('totalScore', '—')}/{result.get('fullMark', '—')} 分（{completed_runs}/{total_runs}）",
                )
            else:
                append_review_job_log(
                    job,
                    f"❌ 「{file_name}」第 {attempt_index} 次失败：{result.get('error') or result.get('reportStatus')}（{completed_runs}/{total_runs}）",
                    "error",
                )

        await asyncio.gather(*[
            run_one(file_index, file_path, attempt_index)
            for file_index, file_path in enumerate(job["files"])
            for attempt_index in range(1, attempts + 1)
        ])

        results.sort(key=lambda item: (int(item.get("fileIndex", 0)), int(item.get("attemptIndex", 0))))
        succeeded = sum(1 for item in results if item.get("success"))
        job["result"] = {
            "jobId": job_id,
            "outputFiles": [],
            "downloadBaseUrl": "",
            "summary": {
                "engine": "skill",
                "skillVersionId": skill_version_id,
                "skillNid": skill_nid,
                "modelName": model_name,
                "attempts": attempts,
                "totalRuns": total_runs,
                "succeededRuns": succeeded,
                "failedRuns": total_runs - succeeded,
                "results": results,
            },
            "scoreTable": build_skill_score_table(results, attempts),
        }
        job["status"] = "completed"
        append_review_job_log(job, f"🎉 Skills 批量测试完成：成功 {succeeded}/{total_runs} 次")
    except asyncio.CancelledError:
        if job.get("status") != "cancelled":
            job["status"] = "cancelled"
            append_review_job_log(job, "⏹️ Skills 批量测试已取消", "warn")
        raise
    except Exception as exc:
        job["status"] = "failed"
        job["error"] = str(exc)
        append_review_job_log(job, f"❌ Skills 批量测试失败：{exc}", "error")
    finally:
        job["updatedAt"] = datetime.now(timezone.utc).isoformat()


async def run_async_skill_review_job(job_id: str, **settings: Any) -> None:
    job = get_review_job(job_id)
    owner_id = job["ownerId"]
    owner_semaphore = get_review_owner_semaphore(owner_id)
    if owner_semaphore.locked():
        append_review_job_log(job, "⏳ 你已有一个批阅任务在执行，本任务在个人队列中等待")
    async with owner_semaphore:
        if job.get("cancelRequested"):
            return
        await execute_async_skill_review_job(job_id, **settings)


@app.get("/")
async def root():
    return {
        "service": "作业批阅API",
        "status": "running",
        "message": "API服务正常运行"
    }


@app.get("/health")
async def health():
    return {"status": "healthy"}


@app.post("/api/review/skill-models")
async def get_skill_review_models(
    authorization: str = Form(...),
    cookie: str = Form(...),
    scene: int = Form(8),
):
    """使用当前智慧树会话读取 Skills 批阅可选模型。"""
    if not authorization.strip() or not cookie.strip():
        raise HTTPException(status_code=400, detail="请填写完整的智慧树认证信息")
    normalized_scene = min(100, max(1, scene))
    try:
        models = await asyncio.to_thread(
            list_correction_skill_models,
            authorization.strip(),
            cookie.strip(),
            scene=normalized_scene,
        )
    except Exception as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    default_model = next(
        (item["code"] for item in models if item.get("isDefault")),
        models[0]["code"],
    )
    return {
        "scene": normalized_scene,
        "models": models,
        "defaultModel": default_model,
    }


@app.post("/api/review/skill-overview")
async def generate_skill_submission_requirement(
    authorization: str = Form(...),
    cookie: str = Form(...),
    skill_version_id: str = Form(...),
    skill_type: str = Form("1", alias="type"),
):
    """根据 Skill 概览生成所有学生共用的可编辑作业要求。"""
    if not authorization.strip() or not cookie.strip():
        raise HTTPException(status_code=400, detail="请填写完整的智慧树认证信息")
    if not skill_version_id.strip():
        raise HTTPException(status_code=400, detail="请填写 Skill Version ID")

    normalized_type = skill_type.strip() or "1"
    try:
        overview = await asyncio.to_thread(
            get_correction_skill_overview,
            skill_version_id.strip(),
            authorization.strip(),
            cookie.strip(),
            skill_type=normalized_type,
        )
        requirement = build_submission_requirement_from_overview(overview)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    data = overview.get("data") or {}
    skill = data.get("skill") or {}
    scoring = data.get("scoring") or {}
    return {
        "requirement": requirement,
        "skillName": skill.get("name"),
        "fullMark": scoring.get("fullMark"),
        "extractionStatus": skill.get("extractionStatus"),
    }


@app.post("/api/review/skill-package")
async def upload_grading_skill_package(
    zip_file: UploadFile = File(..., alias="zipFile"),
    authorization: str = Form(...),
    cookie: str = Form(...),
):
    """校验并上传技能包，随后自动将其改为批阅类型 Skill。"""
    if not authorization.strip() or not cookie.strip():
        raise HTTPException(status_code=400, detail="请填写完整的智慧树认证信息")
    if not zip_file.filename or not zip_file.filename.lower().endswith(".zip"):
        raise HTTPException(status_code=400, detail="请选择 .zip 作业批阅技能包")

    with tempfile.TemporaryDirectory(prefix="grading_skill_upload_") as temp_dir:
        target = Path(temp_dir) / safe_upload_name(zip_file.filename)
        total_bytes = 0
        try:
            with target.open("wb") as output:
                while True:
                    chunk = await zip_file.read(REVIEW_UPLOAD_CHUNK_SIZE)
                    if not chunk:
                        break
                    total_bytes += len(chunk)
                    if total_bytes > MAX_SKILL_PACKAGE_BYTES:
                        raise HTTPException(status_code=413, detail="技能包不能超过 50 MB")
                    output.write(chunk)
        finally:
            await zip_file.close()

        try:
            result = await asyncio.to_thread(
                upload_and_prepare_grading_skill,
                str(target),
                authorization.strip(),
                cookie.strip(),
            )
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc

    return result


@app.post("/api/review/skill-package/generate")
async def generate_upload_grading_skill_package(
    materials: Optional[List[UploadFile]] = File(None),
    material_text: str = Form(""),
    authorization: str = Form(...),
    cookie: str = Form(...),
    llm_api_key: str = Form(""),
    llm_api_url: str = Form(""),
    llm_model: str = Form(""),
):
    """使用 AgentEval LLM 生成 Skill ZIP，并自动上传和切换为批阅类型。"""
    if not authorization.strip() or not cookie.strip():
        raise HTTPException(status_code=400, detail="请填写完整的智慧树认证信息")
    valid_materials = [item for item in (materials or []) if item.filename]
    if not valid_materials and not material_text.strip():
        raise HTTPException(status_code=400, detail="请上传课程材料或填写教师补充说明")
    if len(valid_materials) > MAX_SKILL_MATERIAL_FILES:
        raise HTTPException(status_code=400, detail=f"课程材料最多 {MAX_SKILL_MATERIAL_FILES} 个文件")

    effective_api_key = llm_api_key.strip() or os.getenv("LLM_API_KEY", "")
    effective_api_url = llm_api_url.strip() or os.getenv("LLM_API_URL", "")
    effective_model = llm_model.strip() or os.getenv("LLM_MODEL", "")
    if not effective_api_key or not effective_model:
        raise HTTPException(status_code=400, detail="请先在 AgentEval 全局设置中配置 LLM API Key 和作业批阅模型")

    with tempfile.TemporaryDirectory(prefix="grading_skill_generate_") as temp_dir:
        temp_root = Path(temp_dir)
        material_dir = temp_root / "materials"
        output_dir = temp_root / "output"
        material_dir.mkdir(parents=True, exist_ok=True)
        material_paths: List[Path] = []
        total_bytes = 0
        try:
            for upload in valid_materials:
                target = unique_upload_path(material_dir, upload.filename or "material")
                with target.open("wb") as output:
                    while True:
                        chunk = await upload.read(REVIEW_UPLOAD_CHUNK_SIZE)
                        if not chunk:
                            break
                        total_bytes += len(chunk)
                        if total_bytes > MAX_SKILL_MATERIAL_BYTES:
                            raise HTTPException(status_code=413, detail="课程材料总大小不能超过 80 MB")
                        output.write(chunk)
                material_paths.append(target)
        finally:
            for upload in valid_materials:
                await upload.close()

        try:
            generated = await asyncio.to_thread(
                generate_grading_skill_zip,
                material_paths=material_paths,
                material_text=material_text,
                output_dir=output_dir,
                api_key=effective_api_key,
                api_url=effective_api_url,
                model=effective_model,
            )
            zip_path = Path(generated.pop("zipPath"))
            zip_base64 = await asyncio.to_thread(encode_file_base64, zip_path)

            uploaded = None
            upload_error = None
            try:
                uploaded = await asyncio.to_thread(
                    upload_and_prepare_grading_skill,
                    str(zip_path),
                    authorization.strip(),
                    cookie.strip(),
                )
            except Exception as exc:
                upload_error = str(exc)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc

    return {
        **generated,
        "zipBase64": zip_base64,
        "upload": uploaded,
        "uploadError": upload_error,
        "model": effective_model,
    }


@app.post("/api/review/student-samples/generate")
async def generate_skill_student_samples(
    submission_requirement: str = Form(...),
    assignment_title: str = Form("课程作业"),
    count: int = Form(3),
    llm_api_key: str = Form(""),
    llm_api_url: str = Form(""),
    llm_model: str = Form(""),
):
    """使用 AgentEval LLM 生成匿名多档学生作业 DOCX，供 Skills 直接测试。"""
    if len(submission_requirement.strip()) < 20:
        raise HTTPException(status_code=400, detail="请先生成或填写完整的学生作业要求")
    normalized_count = min(5, max(1, count))
    effective_api_key = llm_api_key.strip() or os.getenv("LLM_API_KEY", "")
    effective_api_url = llm_api_url.strip() or os.getenv("LLM_API_URL", "")
    effective_model = llm_model.strip() or os.getenv("LLM_MODEL", "")
    if not effective_api_key or not effective_model:
        raise HTTPException(status_code=400, detail="请先在 AgentEval 全局设置中配置 LLM API Key 和作业批阅模型")

    with tempfile.TemporaryDirectory(prefix="grading_student_samples_") as temp_dir:
        try:
            files = await asyncio.to_thread(
                generate_student_sample_docx_files,
                assignment_title=assignment_title.strip() or "课程作业",
                submission_requirement=submission_requirement.strip(),
                count=normalized_count,
                output_dir=Path(temp_dir),
                api_key=effective_api_key,
                api_url=effective_api_url,
                model=effective_model,
            )
        except Exception as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc

    return {
        "count": len(files),
        "files": files,
        "model": effective_model,
    }


@app.get("/api/files")
async def download_file(path: str = Query(..., description="文件路径")):
    """下载服务器上的临时文件（用于生成后的批阅流程）"""
    file_path = resolve_temp_file(path)
    if not file_path.is_file():
        raise HTTPException(
            status_code=404, 
            detail=f"文件不存在（可能服务已重启）: {path}"
        )
    return FileResponse(
        file_path,
        filename=file_path.name,
        media_type="application/octet-stream",
    )


@app.get("/api/preview")
async def preview_file(path: str = Query(..., description="文件路径")):
    """预览文件 - 支持 docx/pdf/ppt/pptx"""
    file_path = resolve_temp_file(path)
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail=f"文件不存在: {path}")
    
    ext = file_path.suffix.lower()
    
    if ext in (".ppt", ".pptx"):
        # PPT 文件返回简要信息
        try:
            from pptx import Presentation
            prs = Presentation(str(file_path))
            slide_count = len(prs.slides)
            html = f"<p>📊 PPT 文件: {file_path.name}</p><p>共 {slide_count} 页幻灯片</p><p style='color:#b45309;font-size:0.85em'>PPT 类型作业跳过 LLM 校验</p>"
        except ImportError:
            html = f"<p>📊 PPT 文件: {file_path.name}</p><p style='color:#b45309;font-size:0.85em'>PPT 类型作业跳过 LLM 校验</p>"
        except Exception as e:
            html = f"<p>📊 PPT 文件: {file_path.name}</p><p style='color:red'>预览失败: {str(e)}</p>"
        return {"html": html, "fileName": file_path.name}
    
    if ext == ".pdf":
        html = f"<p>📄 PDF 文件: {file_path.name}</p><p style='color:#6366f1;font-size:0.85em'>PDF 文件将通过云端解析</p>"
        return {"html": html, "fileName": file_path.name}
    
    if ext not in (".docx", ".doc"):
        raise HTTPException(status_code=400, detail="仅支持预览 .docx/.pdf/.ppt/.pptx 文件")
    
    try:
        from docx import Document
        doc = Document(str(file_path))
        paragraphs = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                # 简单样式处理
                if p.style and p.style.name and "Heading" in p.style.name:
                    paragraphs.append(f"<h3>{text}</h3>")
                else:
                    paragraphs.append(f"<p>{text}</p>")
        html = "\n".join(paragraphs) if paragraphs else "<p>文档内容为空</p>"
        return {"html": html, "fileName": file_path.name}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"预览失败: {str(e)}")


@app.get("/test")
async def test():
    """测试环境变量"""
    return {
        "env": {
            "PORT": os.getenv("PORT", "未设置"),
            "LLM_API_KEY": "已设置" if os.getenv("LLM_API_KEY") else "未设置",
            "AUTHORIZATION": "前端传递" if not os.getenv("AUTHORIZATION") else "已设置"
        },
        "generate_script": str(GENERATE_SCRIPT),
        "generate_script_exists": GENERATE_SCRIPT.exists(),
        "review_script": str(REVIEW_SCRIPT),
        "review_script_exists": REVIEW_SCRIPT.exists(),
    }


@app.post("/api/generate")
async def generate_answers(
    file: Optional[UploadFile] = File(None),
    exam_text: Optional[str] = Form(None),
    exam_title: Optional[str] = Form(None),
    authorization: Optional[str] = Form(None),
    cookie: Optional[str] = Form(None),
    instance_nid: Optional[str] = Form(None),
    llm_api_key: Optional[str] = Form(None),
    llm_api_url: Optional[str] = Form(None),
    llm_model: Optional[str] = Form(None),
    levels: Optional[str] = Form(None),
    auto_review: Optional[str] = Form(None),
    custom_prompt: Optional[str] = Form(None),
    custom_levels: Optional[str] = Form(None),
):
    """生成学生答案 - 调用 generate_and_review_service.py"""
    
    # 创建临时目录
    temp_dir = Path(tempfile.mkdtemp(prefix="homework_"))
    output_root = temp_dir / "output"
    output_root.mkdir(exist_ok=True)
    exam_file: Optional[Path] = None
    text_input_file: Optional[Path] = None

    raw_exam_text = exam_text or ""
    if not file and not raw_exam_text.strip():
        raise HTTPException(status_code=400, detail="请上传题卷文件或粘贴题卷文字")

    if file and file.filename:
        exam_file = temp_dir / file.filename
        content = await file.read()
        exam_file.write_bytes(content)
    elif raw_exam_text.strip():
        safe_title = (exam_title or "作业").strip() or "作业"
        safe_title = "".join(ch if ch.isalnum() or ch in "-_." or "\u4e00" <= ch <= "\u9fff" else "_" for ch in safe_title)[:40]
        text_input_file = temp_dir / f"{safe_title or '作业'}.txt"
        text_input_file.write_text(raw_exam_text, encoding="utf-8")
    
    # 解析等级 - 前端发送的是字符串数组 ["优秀的回答", "良好的回答", ...]
    levels_list = ["优秀的回答", "良好的回答", "中等的回答", "合格的回答", "较差的回答"]
    if levels:
        try:
            parsed = json.loads(levels)
            if isinstance(parsed, list) and len(parsed) > 0:
                levels_list = parsed
        except json.JSONDecodeError:
            pass
    
    # 创建环境变量 - 始终设置认证变量（子脚本会检查这些变量来决定是否加载.env）
    env = os.environ.copy()
    env["PYTHONUNBUFFERED"] = "1"
    env["AUTHORIZATION"] = authorization or ""
    env["COOKIE"] = cookie or ""
    env["INSTANCE_NID"] = instance_nid or ""
    env["LLM_API_KEY"] = llm_api_key or os.getenv("LLM_API_KEY", "")
    env["LLM_API_URL"] = llm_api_url or os.getenv("LLM_API_URL", "")
    env["LLM_MODEL"] = llm_model or os.getenv("LLM_MODEL", "")
    if custom_prompt:
        env["CUSTOM_PROMPT"] = custom_prompt
    if custom_levels:
        env["CUSTOM_LEVELS"] = custom_levels
    
    # 构建命令行参数 - 与前端本地模式一致
    cmd = [
        sys.executable, "-u",
        str(GENERATE_SCRIPT),
        "--output-root", str(output_root),
        "--levels", *levels_list,
        "--llm-api-key", llm_api_key or os.getenv("LLM_API_KEY", ""),
        "--llm-api-url", llm_api_url or os.getenv("LLM_API_URL", ""),
        "--llm-model", llm_model or os.getenv("LLM_MODEL", ""),
    ]
    if exam_file is not None:
        cmd[3:3] = ["--input", str(exam_file)]
    elif text_input_file is not None:
        cmd[3:3] = ["--input-text-file", str(text_input_file)]
        if exam_title and exam_title.strip():
            cmd[5:5] = ["--input-title", exam_title.strip()]
    
    async def event_stream():
        """SSE流式响应 - 读取子进程的JSON行协议输出"""
        try:
            process = await asyncio.create_subprocess_exec(
                *cmd,
                env=env,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
                cwd=str(SCRIPT_DIR),
                limit=10 * 1024 * 1024,  # 10MB行缓冲，防止大JSON超限
            )
            
            # 读取stdout（JSON行协议），带心跳保活防止Railway空闲超时
            while True:
                try:
                    line = await asyncio.wait_for(process.stdout.readline(), timeout=15)
                except asyncio.TimeoutError:
                    # 子进程无输出时发送SSE心跳注释，防止Railway 60秒空闲断连
                    yield ": heartbeat\n\n"
                    continue
                if not line:
                    break
                msg = line.decode().strip()
                if not msg:
                    continue
                # generate_and_review_service.py 输出JSON行协议
                # 直接转发为SSE
                try:
                    data = json.loads(msg)
                    yield f'data: {json.dumps(data, ensure_ascii=False)}\n\n'
                except json.JSONDecodeError:
                    # 非JSON行作为日志
                    yield f'data: {json.dumps({"type": "log", "message": msg}, ensure_ascii=False)}\n\n'
            
            await process.wait()
            
            if process.returncode != 0:
                stderr = await process.stderr.read()
                err_msg = stderr.decode().strip()
                if err_msg:
                    yield f'data: {json.dumps({"type": "error", "message": err_msg}, ensure_ascii=False)}\n\n'
                
        except Exception as e:
            yield f'data: {json.dumps({"type": "error", "message": str(e)}, ensure_ascii=False)}\n\n'
        finally:
            yield f'data: {json.dumps({"type": "done"})}\n\n'
    
    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )


@app.post("/api/review/jobs", status_code=201)
async def create_review_job(review_user_id: str = Depends(require_review_user)):
    """Create an upload session for a long-running review job."""
    job_id = uuid.uuid4().hex
    job_dir = REVIEW_JOBS_ROOT / job_id
    upload_dir = job_dir / "uploads"
    output_root = job_dir / "output"
    upload_dir.mkdir(parents=True, exist_ok=False)
    output_root.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone.utc).isoformat()
    REVIEW_JOBS[job_id] = {
        "jobId": job_id,
        "ownerId": review_user_id,
        "status": "uploading",
        "files": [],
        "uploadedNames": [],
        "uploadBatches": {},
        "chunkUploads": {},
        "uploadDir": str(upload_dir),
        "outputRoot": str(output_root),
        "logs": [],
        "result": None,
        "error": None,
        "cancelRequested": False,
        "createdAt": now,
        "updatedAt": now,
    }
    return {
        "jobId": job_id,
        "status": "uploading",
        "maxFiles": MAX_REVIEW_FILES,
        "maxConcurrency": MAX_REVIEW_CONCURRENCY,
    }


@app.post("/api/review/jobs/{job_id}/files")
async def upload_review_job_files(
    job_id: str,
    files: List[UploadFile] = File(...),
    upload_batch_id: str = Form(...),
    file_keys: Optional[str] = Form(None),
    review_user_id: str = Depends(require_review_user),
):
    """Upload one idempotent chunk of files into an asynchronous job."""
    job = get_review_job(job_id, review_user_id)
    if job["status"] != "uploading":
        raise HTTPException(status_code=409, detail="当前任务已结束上传阶段")

    batch_id = upload_batch_id.strip()
    if not batch_id:
        raise HTTPException(status_code=400, detail="upload_batch_id 不能为空")
    cached = job["uploadBatches"].get(batch_id)
    if cached:
        return cached

    valid_files = [item for item in files if item.filename]
    if not valid_files:
        raise HTTPException(status_code=400, detail="请至少上传一个作业文件")
    pending_chunks = sum(1 for item in job["chunkUploads"].values() if not item.get("complete"))
    if len(job["files"]) + pending_chunks + len(valid_files) > MAX_REVIEW_FILES:
        raise HTTPException(status_code=400, detail=f"每个批阅任务最多 {MAX_REVIEW_FILES} 份文件")

    parsed_keys: List[str] = []
    if file_keys:
        try:
            value = json.loads(file_keys)
            if isinstance(value, list):
                parsed_keys = [str(item) for item in value]
        except json.JSONDecodeError:
            pass

    upload_dir = Path(job["uploadDir"])
    uploaded = []
    for index, upload in enumerate(valid_files):
        target = unique_upload_path(upload_dir, upload.filename or "file")
        with target.open("wb") as output:
            while True:
                chunk = await upload.read(REVIEW_UPLOAD_CHUNK_SIZE)
                if not chunk:
                    break
                output.write(chunk)
        await upload.close()
        job["files"].append(str(target))
        job["uploadedNames"].append(target.name)
        uploaded.append({
            "clientKey": parsed_keys[index] if index < len(parsed_keys) else "",
            "originalName": upload.filename,
            "storedName": target.name,
        })

    response = {
        "jobId": job_id,
        "uploaded": uploaded,
        "uploadedCount": len(job["files"]),
        "maxFiles": MAX_REVIEW_FILES,
    }
    job["uploadBatches"][batch_id] = response
    append_review_job_log(job, f"📦 已上传 {len(job['files'])}/{MAX_REVIEW_FILES} 份文件")
    return response


@app.post("/api/review/jobs/{job_id}/chunks")
async def upload_review_job_chunk(
    job_id: str,
    file: UploadFile = File(...),
    upload_id: str = Form(...),
    client_key: str = Form(""),
    original_name: str = Form(...),
    chunk_index: int = Form(...),
    total_chunks: int = Form(...),
    review_user_id: str = Depends(require_review_user),
):
    """Append one idempotent chunk for a large file."""
    job = get_review_job(job_id, review_user_id)
    if job["status"] != "uploading":
        raise HTTPException(status_code=409, detail="当前任务已结束上传阶段")
    if not upload_id.strip() or total_chunks < 1 or chunk_index < 0 or chunk_index >= total_chunks:
        raise HTTPException(status_code=400, detail="分片参数不正确")

    chunk_uploads = job["chunkUploads"]
    state = chunk_uploads.get(upload_id)
    if state is None:
        pending_count = sum(1 for item in chunk_uploads.values() if not item.get("complete"))
        if len(job["files"]) + pending_count >= MAX_REVIEW_FILES:
            raise HTTPException(status_code=400, detail=f"每个批阅任务最多 {MAX_REVIEW_FILES} 份文件")
        target = unique_upload_path(Path(job["uploadDir"]), original_name)
        state = {
            "target": str(target),
            "part": str(target.with_name(f".{target.name}.{upload_id}.part")),
            "nextChunk": 0,
            "totalChunks": total_chunks,
            "clientKey": client_key,
            "originalName": original_name,
            "complete": False,
        }
        chunk_uploads[upload_id] = state

    if state["complete"]:
        return {
            "jobId": job_id,
            "complete": True,
            "uploadedCount": len(job["files"]),
            "uploaded": {
                "clientKey": state["clientKey"],
                "originalName": state["originalName"],
                "storedName": Path(state["target"]).name,
            },
        }
    if total_chunks != state["totalChunks"]:
        raise HTTPException(status_code=409, detail="分片总数与已建立的上传会话不一致")
    if chunk_index < state["nextChunk"]:
        return {
            "jobId": job_id,
            "complete": False,
            "nextChunk": state["nextChunk"],
            "uploadedCount": len(job["files"]),
        }
    if chunk_index != state["nextChunk"]:
        raise HTTPException(status_code=409, detail=f"请先上传第 {state['nextChunk'] + 1} 个分片")

    content = await file.read()
    await file.close()
    with Path(state["part"]).open("ab") as output:
        output.write(content)
    state["nextChunk"] += 1

    if state["nextChunk"] == state["totalChunks"]:
        part_path = Path(state["part"])
        target_path = Path(state["target"])
        part_path.replace(target_path)
        state["complete"] = True
        job["files"].append(str(target_path))
        job["uploadedNames"].append(target_path.name)
        append_review_job_log(job, f"📦 已上传 {len(job['files'])}/{MAX_REVIEW_FILES} 份文件")

    return {
        "jobId": job_id,
        "complete": state["complete"],
        "nextChunk": state["nextChunk"],
        "uploadedCount": len(job["files"]),
        "uploaded": {
            "clientKey": state["clientKey"],
            "originalName": state["originalName"],
            "storedName": Path(state["target"]).name,
        } if state["complete"] else None,
    }


@app.post("/api/review/jobs/{job_id}/start", status_code=202)
async def start_review_job(
    job_id: str,
    authorization: str = Form(...),
    cookie: str = Form(...),
    instance_nid: str = Form(...),
    attempts: int = Form(5),
    output_format: str = Form("json"),
    max_concurrency: int = Form(DEFAULT_REVIEW_CONCURRENCY),
    local_parse: bool = Form(False),
    llm_api_key: Optional[str] = Form(None),
    llm_api_url: Optional[str] = Form(None),
    llm_model: Optional[str] = Form(None),
    skip_llm_files: Optional[str] = Form(None),
    file_groups: Optional[str] = Form(None),
    review_user_id: str = Depends(require_review_user),
):
    """Start the worker and return immediately; progress is read by polling."""
    job = get_review_job(job_id, review_user_id)
    if job["status"] in {"queued", "running", "completed"}:
        return {"jobId": job_id, "status": job["status"]}
    if job["status"] == "cancelled":
        raise HTTPException(status_code=409, detail="任务已取消")
    if job["status"] == "failed":
        raise HTTPException(status_code=409, detail=job.get("error") or "任务已失败")
    if not job["files"]:
        raise HTTPException(status_code=400, detail="请先上传作业文件")
    if any(not item.get("complete") for item in job["chunkUploads"].values()):
        raise HTTPException(status_code=409, detail="仍有大文件分片未上传完成")
    if not authorization.strip() or not cookie.strip() or not instance_nid.strip():
        raise HTTPException(status_code=400, detail="请填写完整的智慧树认证信息")
    if output_format not in {"json", "pdf"}:
        raise HTTPException(status_code=400, detail="不支持的输出格式")

    concurrency = clamp_review_concurrency(max_concurrency)
    job["status"] = "queued"
    job["configuredConcurrency"] = concurrency
    append_review_job_log(job, f"✅ {len(job['files'])} 份文件已就绪，进入批阅队列")
    task = asyncio.create_task(run_async_review_job(
        job_id,
        authorization=authorization.strip(),
        cookie=cookie.strip(),
        instance_nid=instance_nid.strip(),
        attempts=max(1, attempts),
        output_format=output_format,
        max_concurrency=concurrency,
        local_parse=local_parse,
        llm_api_key=(llm_api_key or "").strip(),
        llm_api_url=(llm_api_url or "").strip(),
        llm_model=(llm_model or "").strip(),
        skip_llm_files=skip_llm_files,
        file_groups=file_groups,
    ))
    REVIEW_JOB_TASKS.add(task)
    job["_task"] = task

    def clear_task(completed_task: asyncio.Task) -> None:
        REVIEW_JOB_TASKS.discard(completed_task)
        if job.get("_task") is completed_task:
            job.pop("_task", None)

    task.add_done_callback(clear_task)
    return {"jobId": job_id, "status": "queued", "maxConcurrency": concurrency}


@app.post("/api/review/jobs/{job_id}/start-skill", status_code=202)
async def start_skill_review_job(
    job_id: str,
    authorization: str = Form(...),
    cookie: str = Form(...),
    skill_version_id: str = Form(...),
    skill_nid: str = Form(""),
    submission_requirement: str = Form(...),
    student_submission: str = Form("见附件"),
    model_name: str = Form("claude-opus-4-8"),
    attempts: int = Form(1),
    max_concurrency: int = Form(3),
    poll_interval_seconds: int = Form(5),
    review_user_id: str = Depends(require_review_user),
):
    """启动多份学生作业的 Skills 批量测试。"""
    job = get_review_job(job_id, review_user_id)
    if job["status"] in {"queued", "running", "completed"}:
        return {"jobId": job_id, "status": job["status"]}
    if job["status"] == "cancelled":
        raise HTTPException(status_code=409, detail="任务已取消")
    if job["status"] == "failed":
        raise HTTPException(status_code=409, detail=job.get("error") or "任务已失败")
    if not job["files"]:
        raise HTTPException(status_code=400, detail="请先上传学生作业文件")
    if any(not item.get("complete") for item in job["chunkUploads"].values()):
        raise HTTPException(status_code=409, detail="仍有大文件分片未上传完成")
    if not authorization.strip() or not cookie.strip():
        raise HTTPException(status_code=400, detail="请填写完整的智慧树认证信息")
    if not skill_version_id.strip():
        raise HTTPException(status_code=400, detail="请填写 Skill Version ID")
    if not submission_requirement.strip():
        raise HTTPException(status_code=400, detail="请填写所有学生共用的作业要求")
    if not model_name.strip():
        raise HTTPException(status_code=400, detail="请填写 Skills 批阅模型")

    normalized_attempts = min(20, max(1, attempts))
    concurrency = clamp_review_concurrency(max_concurrency)
    poll_interval = min(30, max(2, poll_interval_seconds))
    job["status"] = "queued"
    job["engine"] = "skill"
    job["configuredConcurrency"] = concurrency
    append_review_job_log(
        job,
        f"✅ {len(job['files'])} 份文件已就绪，Skills 批量测试进入队列",
    )
    task = asyncio.create_task(run_async_skill_review_job(
        job_id,
        authorization=authorization.strip(),
        cookie=cookie.strip(),
        skill_version_id=skill_version_id.strip(),
        skill_nid=skill_nid.strip(),
        submission_requirement=submission_requirement.strip(),
        student_submission=student_submission.strip() or "见附件",
        model_name=model_name.strip(),
        attempts=normalized_attempts,
        max_concurrency=concurrency,
        poll_interval_seconds=poll_interval,
    ))
    REVIEW_JOB_TASKS.add(task)
    job["_task"] = task

    def clear_task(completed_task: asyncio.Task) -> None:
        REVIEW_JOB_TASKS.discard(completed_task)
        if job.get("_task") is completed_task:
            job.pop("_task", None)

    task.add_done_callback(clear_task)
    return {
        "jobId": job_id,
        "status": "queued",
        "engine": "skill",
        "attempts": normalized_attempts,
        "maxConcurrency": concurrency,
    }


@app.get("/api/review/jobs/active")
async def get_active_review_jobs(
    review_user_id: str = Depends(require_review_user),
):
    """List this user's running and queued jobs so stale page state can recover."""
    active_jobs = [
        job
        for job in REVIEW_JOBS.values()
        if job.get("ownerId") == review_user_id
        and job.get("status") in {"running", "queued"}
    ]
    active_jobs.sort(
        key=lambda item: (
            0 if item.get("status") == "running" else 1,
            str(item.get("createdAt") or ""),
        )
    )
    return {
        "jobs": [review_job_summary(job) for job in active_jobs],
        "runningCount": sum(1 for job in active_jobs if job.get("status") == "running"),
        "queuedCount": sum(1 for job in active_jobs if job.get("status") == "queued"),
    }


@app.delete("/api/review/jobs/active")
async def cancel_active_review_job(
    review_user_id: str = Depends(require_review_user),
):
    """Cancel the oldest running job while leaving later queued jobs intact."""
    running_jobs = [
        job
        for job in REVIEW_JOBS.values()
        if job.get("ownerId") == review_user_id and job.get("status") == "running"
    ]
    if not running_jobs:
        raise HTTPException(status_code=404, detail="当前账号没有正在执行的批阅任务")

    target = min(running_jobs, key=lambda item: str(item.get("createdAt") or ""))
    await cancel_review_job_state(target, "⏹️ 用户手动结束了占用中的批阅任务")
    return {
        "jobId": target["jobId"],
        "status": "cancelled",
        "engine": target.get("engine") or "traditional",
        "message": "正在执行的批阅任务已结束，个人队列已释放",
    }


@app.delete("/api/review/jobs/{job_id}")
async def cancel_review_job(
    job_id: str,
    review_user_id: str = Depends(require_review_user),
):
    """Cancel an upload, queued job, or running review process."""
    job = get_review_job(job_id, review_user_id)
    if job["status"] == "cancelled":
        return {"jobId": job_id, "status": "cancelled"}
    if job["status"] in {"completed", "failed"}:
        raise HTTPException(status_code=409, detail=f"任务已{'完成' if job['status'] == 'completed' else '失败'}")

    await cancel_review_job_state(job, "⏹️ 用户已取消本次批阅")
    return {"jobId": job_id, "status": "cancelled"}


@app.get("/api/review/jobs/{job_id}/artifacts")
async def download_review_job_artifact(
    job_id: str,
    file: str = Query(..., min_length=1),
    review_user_id: str = Depends(require_review_user),
):
    """Download one completed result after checking the job owner."""
    job = get_review_job(job_id, review_user_id)
    if job["status"] != "completed":
        raise HTTPException(status_code=409, detail="批阅任务尚未完成")
    artifact = resolve_review_job_artifact(job, file)
    return FileResponse(path=str(artifact), filename=artifact.name)


@app.get("/api/review/jobs/{job_id}")
async def get_review_job_status(
    job_id: str,
    cursor: int = Query(0, ge=0),
    log_limit: int = Query(250, ge=1, le=1000),
    review_user_id: str = Depends(require_review_user),
):
    """Return a compact status snapshot and logs after the supplied cursor."""
    job = get_review_job(job_id, review_user_id)
    end = min(len(job["logs"]), cursor + log_limit)
    response: Dict[str, Any] = {
        "jobId": job_id,
        "status": job["status"],
        "uploadedCount": len(job["files"]),
        "configuredConcurrency": job.get("configuredConcurrency"),
        "logs": job["logs"][cursor:end],
        "nextCursor": end,
        "hasMoreLogs": end < len(job["logs"]),
        "error": job.get("error"),
        "createdAt": job["createdAt"],
        "updatedAt": job["updatedAt"],
    }
    if job["status"] == "completed":
        response["result"] = job["result"]
    return response


@app.post("/api/review")
async def review_answers(
    files: List[UploadFile] = File(None),
    server_paths: Optional[str] = Form(None),
    authorization: Optional[str] = Form(None),
    cookie: Optional[str] = Form(None),
    instance_nid: Optional[str] = Form(None),
    attempts: int = Form(5),
    max_workers: int = Form(3),
    output_format: str = Form("json"),
    max_concurrency: int = Form(DEFAULT_REVIEW_CONCURRENCY),
    local_parse: bool = Form(False),
    llm_api_key: Optional[str] = Form(None),
    llm_api_url: Optional[str] = Form(None),
    llm_model: Optional[str] = Form(None),
    skip_llm_files: Optional[str] = Form(None),
    file_groups: Optional[str] = Form(None),
):
    """批阅学生答案 - 调用 review_service.py"""
    
    # 创建临时目录
    temp_dir = Path(tempfile.mkdtemp(prefix="homework_"))
    output_root = temp_dir / "output"
    output_root.mkdir(exist_ok=True)
    student_files = []

    incoming_file_count = len([item for item in (files or []) if item.filename])
    if incoming_file_count > MAX_REVIEW_FILES:
        raise HTTPException(status_code=400, detail=f"每个批阅任务最多 {MAX_REVIEW_FILES} 份文件")
    
    # 保存上传的文件
    if files:
        for f in files:
            if f.filename:
                target = temp_dir / f.filename
                content = await f.read()
                target.write_bytes(content)
                student_files.append(str(target))
    
    # 或使用服务器路径
    if server_paths:
        try:
            paths = json.loads(server_paths)
            if len(paths) + incoming_file_count > MAX_REVIEW_FILES:
                raise HTTPException(status_code=400, detail=f"每个批阅任务最多 {MAX_REVIEW_FILES} 份文件")
            student_files.extend(paths)
        except json.JSONDecodeError:
            pass
    
    if not student_files:
        raise HTTPException(status_code=400, detail="请提供至少一个学生答案文件")
    
    # 创建环境变量 - 始终设置认证变量（子脚本会检查这些变量来决定是否加载.env）
    env = os.environ.copy()
    env["PYTHONUNBUFFERED"] = "1"
    env["AUTHORIZATION"] = authorization or ""
    env["COOKIE"] = cookie or ""
    env["INSTANCE_NID"] = instance_nid or ""
    env["LLM_API_KEY"] = llm_api_key or os.getenv("LLM_API_KEY", "")
    env["LLM_API_URL"] = llm_api_url or os.getenv("LLM_API_URL", "")
    env["LLM_MODEL"] = llm_model or os.getenv("LLM_MODEL", "")
    
    # 构建命令 - 与前端本地模式一致，调用 review_service.py
    cmd = [
        sys.executable, "-u",
        str(REVIEW_SCRIPT),
        "--inputs", json.dumps(student_files),
        "--attempts", str(max(1, attempts)),
        "--output-format", output_format,
        "--output-root", str(output_root),
        "--max-concurrency", str(clamp_review_concurrency(max_concurrency)),
    ]
    if local_parse:
        cmd.append("--local-parse")
    if skip_llm_files:
        cmd.extend(["--skip-llm-files", skip_llm_files])
    if file_groups:
        cmd.extend(["--file-groups", file_groups])
    
    async def event_stream():
        """SSE流式响应"""
        try:
            process = await asyncio.create_subprocess_exec(
                *cmd,
                env=env,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
                cwd=str(SCRIPT_DIR),
                limit=10 * 1024 * 1024,  # 10MB行缓冲，防止__RESULT__大JSON超限
            )
            
            # 读取stdout，带心跳保活防止Railway空闲超时
            while True:
                try:
                    line = await asyncio.wait_for(process.stdout.readline(), timeout=15)
                except asyncio.TimeoutError:
                    # 子进程无输出时发送SSE心跳注释，防止Railway 60秒空闲断连
                    yield ": heartbeat\n\n"
                    continue
                if not line:
                    break
                msg = line.decode().strip()
                if not msg:
                    continue
                
                # 识别 __RESULT__ 标记（review_service.py 的最终输出）
                if msg.startswith("__RESULT__"):
                    try:
                        payload = json.loads(msg[len("__RESULT__"):])
                        # 将相对路径转为绝对路径，前端用 /api/files?path= 下载
                        rel_files = payload.get("output_files", [])
                        abs_files = [str(output_root / f) for f in rel_files]
                        # 转换为前端期望的 "complete" 事件格式
                        complete_event = {
                            "type": "complete",
                            "jobId": "",
                            "outputFiles": abs_files,
                            "summary": payload.get("result", {}),
                            "scoreTable": payload.get("score_table", None),
                            "downloadBaseUrl": "/api/homework-review/download",
                        }
                        yield f'data: {json.dumps(complete_event, ensure_ascii=False)}\n\n'
                    except json.JSONDecodeError:
                        yield f'data: {json.dumps({"type": "log", "message": msg}, ensure_ascii=False)}\n\n'
                    continue
                
                try:
                    data = json.loads(msg)
                    yield f'data: {json.dumps(data, ensure_ascii=False)}\n\n'
                except json.JSONDecodeError:
                    yield f'data: {json.dumps({"type": "log", "message": msg}, ensure_ascii=False)}\n\n'
            
            await process.wait()
            
            if process.returncode != 0:
                stderr = await process.stderr.read()
                err_msg = stderr.decode().strip()
                if err_msg:
                    yield f'data: {json.dumps({"type": "error", "message": err_msg}, ensure_ascii=False)}\n\n'
                
        except Exception as e:
            yield f'data: {json.dumps({"type": "error", "message": str(e)}, ensure_ascii=False)}\n\n'
        finally:
            yield f'data: {json.dumps({"type": "done"})}\n\n'
    
    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
        }
    )


if __name__ == "__main__":
    import uvicorn
    port = int(os.getenv("PORT", 8000))
    uvicorn.run("main:app", host="0.0.0.0", port=port)
